from __future__ import annotations

import io
import json
import uuid

import pytest
from django.test import Client
from django.utils import timezone
from openpyxl import Workbook, load_workbook
from PIL import Image

from apps.inventory.alerts import (
    active_stock_alerts,
    update_inventory_low_stock_threshold,
    update_product_low_stock_threshold,
)
from apps.inventory.bulk import (
    analyse_import_job,
    build_import_template,
    confirm_inventory_import,
    inventory_export_download_url,
    mapping_from_import_headers,
    preview_inventory_import,
    process_export_job,
    process_import_job,
    retry_inventory_export,
    retry_inventory_import,
    start_inventory_export,
    start_inventory_import,
)
from apps.inventory.media import (
    attach_product_media,
    remove_product_media,
    set_primary_product_media,
)
from apps.inventory.models import (
    CustomFieldDefinition,
    InventoryAlert,
    InventoryExport,
    InventoryImport,
    Product,
    ProductMediaAttachment,
    StockMovement,
)
from apps.inventory.selectors import inventory_summary
from apps.inventory.services import create_custom_field, create_product, record_stock_movement
from apps.media_assets.models import MediaAsset
from apps.media_assets.services import complete_upload, prepare_upload
from apps.media_assets.storage import fake_object_storage
from apps.organisations.selectors import TenantContext, resolve_tenant_context
from apps.organisations.services import create_organisation_for_owner
from apps.users.models import User
from tenda.errors import DomainError

pytestmark = pytest.mark.django_db(transaction=True)


def context_for(email: str) -> TenantContext:
    user = User.objects.create_user(
        email=email,
        password="Correct-Horse-Battery-42",
        email_verified_at=timezone.now(),
    )
    create_organisation_for_owner(owner=user, name=f"Negocio {email}")
    return resolve_tenant_context(user)


def rgb_image_bytes() -> bytes:
    buffer = io.BytesIO()
    Image.new("RGB", (48, 48), (20, 80, 40)).save(buffer, format="PNG")
    return buffer.getvalue()


def uploaded_asset(
    context: TenantContext,
    *,
    purpose: str,
    name: str,
    content_type: str,
    content: bytes,
) -> MediaAsset:
    prepared = prepare_upload(
        context=context,
        purpose=purpose,
        original_name=name,
        content_type=content_type,
        size=len(content),
    )
    fake_object_storage.write_bytes(
        key=prepared.asset.object_key,
        content=content,
        content_type=content_type,
    )
    return complete_upload(context=context, public_id=prepared.asset.public_id)


def test_product_media_is_private_tenant_safe_and_has_one_primary() -> None:
    fake_object_storage.clear()
    context = context_for("media-owner@example.com")
    foreign = context_for("media-foreign@example.com")
    product = create_product(context=context, name="Vela").product
    first = uploaded_asset(
        context,
        purpose=MediaAsset.Purpose.PRODUCT_IMAGE,
        name="frente.png",
        content_type="image/png",
        content=rgb_image_bytes(),
    )
    second = uploaded_asset(
        context,
        purpose=MediaAsset.Purpose.PRODUCT_IMAGE,
        name="detalle.png",
        content_type="image/png",
        content=rgb_image_bytes(),
    )

    attachment = attach_product_media(
        context=context,
        product_id=product.public_id,
        asset_id=first.public_id,
    )
    attach_product_media(
        context=context,
        product_id=product.public_id,
        asset_id=second.public_id,
    )
    product.refresh_from_db()
    assert product.primary_image_id == first.id
    assert attachment.asset.organisation_id == context.organisation.id
    assert ProductMediaAttachment.objects.filter(product=product).count() == 2

    set_primary_product_media(
        context=context,
        product_id=product.public_id,
        asset_id=second.public_id,
    )
    product.refresh_from_db()
    assert product.primary_image_id == second.id

    with pytest.raises(DomainError):
        attach_product_media(
            context=foreign,
            product_id=product.public_id,
            asset_id=first.public_id,
        )

    remove_product_media(
        context=context,
        product_id=product.public_id,
        asset_id=second.public_id,
    )
    product.refresh_from_db()
    assert product.primary_image_id == first.id
    assert not MediaAsset.objects.filter(pk=second.pk).exists()


def test_product_media_graphql_urls_use_content_endpoint() -> None:
    fake_object_storage.clear()
    email = "media-urls@example.com"
    context = context_for(email)
    product = create_product(context=context, name="Vela").product
    asset = uploaded_asset(
        context,
        purpose=MediaAsset.Purpose.PRODUCT_IMAGE,
        name="frente.png",
        content_type="image/png",
        content=rgb_image_bytes(),
    )
    attach_product_media(
        context=context,
        product_id=product.public_id,
        asset_id=asset.public_id,
    )
    client = Client()
    login = client.post(
        "/api/v1/auth/login",
        data=json.dumps({"email": email, "password": "Correct-Horse-Battery-42"}),
        content_type="application/json",
    )
    assert login.status_code == 200
    response = client.post(
        "/graphql/",
        data=json.dumps(
            {
                "query": """
                  query ($id: ID!) {
                    product(id: $id) {
                      media { url thumbnailUrl mediumUrl largeUrl }
                    }
                  }
                """,
                "variables": {"id": str(product.public_id)},
            }
        ),
        content_type="application/json",
    )
    body = response.json()
    assert "errors" not in body, body
    media = body["data"]["product"]["media"][0]
    prefix = f"/api/v1/media/{asset.public_id}/content"
    assert prefix in media["url"]
    assert media["url"].endswith("variant=medium")
    assert media["thumbnailUrl"].endswith("variant=thumbnail")
    assert media["mediumUrl"].endswith("variant=medium")
    assert media["largeUrl"].endswith("variant=large")
    assert "r2.invalid" not in media["url"]
    content = client.get(f"{prefix}?variant=medium")
    assert content.status_code == 200


def test_import_keeps_valid_rows_reports_errors_and_replays_without_duplicates() -> None:
    fake_object_storage.clear()
    context = context_for("import-owner@example.com")
    create_custom_field(
        context=context,
        label="Aroma",
        key="aroma",
        field_type=CustomFieldDefinition.FieldType.SHORT_TEXT,
    )
    source = uploaded_asset(
        context,
        purpose=MediaAsset.Purpose.IMPORT_FILE,
        name="productos.csv",
        content_type="text/csv",
        content=(
            "Nombre;Cantidad;Aroma;Precio\n"
            "Vela lavanda;5;Lavanda;4500\n"
            "Vela inválida;-3;Canela;3000\n"
        ).encode(),
    )

    job = start_inventory_import(context=context, asset_id=source.public_id)
    analysed = analyse_import_job(job.public_id)
    assert analysed.status == InventoryImport.Status.AWAITING_MAPPING
    assert analysed.headers == ["Nombre", "Cantidad", "Aroma", "Precio"]

    previewed = preview_inventory_import(
        context=context,
        import_id=job.public_id,
        mapping={
            "name": "Nombre",
            "initialQuantity": "Cantidad",
            "salePrice": "Precio",
            "extraAttributes.aroma": "Aroma",
        },
    )
    assert previewed.row_errors[0]["row"] == 3

    queued, replayed = confirm_inventory_import(
        context=context,
        import_id=job.public_id,
        idempotency_key="confirm-products-1",
    )
    assert queued.status == InventoryImport.Status.QUEUED
    assert replayed is False
    completed = process_import_job(job.public_id)
    assert completed.status == InventoryImport.Status.COMPLETED_WITH_ERRORS
    assert completed.created_count == 1
    assert completed.error_count == 1
    assert completed.row_errors[0]["row"] == 3
    assert completed.report_asset is not None
    assert Product.objects.filter(inventory=context.inventory, name="Vela lavanda").count() == 1
    assert StockMovement.objects.filter(product__name="Vela lavanda").count() == 1

    retry_inventory_import(context=context, import_id=job.public_id)
    replay = process_import_job(job.public_id)
    assert replay.created_count == 1
    assert Product.objects.filter(inventory=context.inventory, name="Vela lavanda").count() == 1
    assert StockMovement.objects.filter(product__name="Vela lavanda").count() == 1

    confirmed_again, was_replayed = confirm_inventory_import(
        context=context,
        import_id=job.public_id,
        idempotency_key="confirm-products-1",
    )
    assert confirmed_again.pk == job.pk
    assert was_replayed is True


def test_import_rejects_dynamic_keys_not_declared_in_inventory() -> None:
    fake_object_storage.clear()
    context = context_for("import-schema@example.com")
    source = uploaded_asset(
        context,
        purpose=MediaAsset.Purpose.IMPORT_FILE,
        name="productos.csv",
        content_type="text/csv",
        content=b"Nombre,Secreto\nVela,valor\n",
    )
    job = start_inventory_import(context=context, asset_id=source.public_id)
    analyse_import_job(job.public_id)

    with pytest.raises(DomainError) as failure:
        preview_inventory_import(
            context=context,
            import_id=job.public_id,
            mapping={"name": "Nombre", "extraAttributes.no_declarada": "Secreto"},
        )

    assert failure.value.code == "INVALID_IMPORT_MAPPING"
    assert "mapping.extraAttributes.no_declarada" in failure.value.field_errors


def test_xlsx_import_and_export_preserve_integral_inventory_values() -> None:
    fake_object_storage.clear()
    context = context_for("xlsx-owner@example.com")
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.append(["Nombre", "Cantidad", "Precio"])
    sheet.append(["Tazón", 7, 12990])
    buffer = io.BytesIO()
    workbook.save(buffer)
    source = uploaded_asset(
        context,
        purpose=MediaAsset.Purpose.IMPORT_FILE,
        name="productos.xlsx",
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        content=buffer.getvalue(),
    )
    job = start_inventory_import(context=context, asset_id=source.public_id)
    analyse_import_job(job.public_id)
    preview_inventory_import(
        context=context,
        import_id=job.public_id,
        mapping={
            "name": "Nombre",
            "initialQuantity": "Cantidad",
            "salePrice": "Precio",
        },
    )
    confirm_inventory_import(
        context=context,
        import_id=job.public_id,
        idempotency_key="confirm-xlsx-1",
    )
    completed_import = process_import_job(job.public_id)
    assert completed_import.status == InventoryImport.Status.SUCCEEDED

    export, _replayed = start_inventory_export(
        context=context,
        file_format=InventoryExport.FileFormat.XLSX,
        filters={},
        idempotency_key="export-xlsx-1",
    )
    completed_export = process_export_job(export.public_id)
    assert completed_export.file_asset is not None
    exported = fake_object_storage.read_bytes(key=completed_export.file_asset.object_key)
    exported_workbook = load_workbook(io.BytesIO(exported), read_only=True, data_only=True)
    exported_sheet = exported_workbook.active
    assert exported_sheet is not None
    rows = list(exported_sheet.iter_rows(values_only=True))
    assert rows[1][0] == "Tazón"
    assert rows[1][2] == 7
    assert rows[1][6] == "12990"


def test_import_template_roundtrip_creates_products_from_canonical_headers() -> None:
    fake_object_storage.clear()
    context = context_for("template-owner@example.com")
    create_custom_field(
        context=context,
        label="Aroma",
        key="aroma",
        field_type=CustomFieldDefinition.FieldType.SHORT_TEXT,
    )
    _file_name, _content_type, content = build_import_template(context)
    workbook = load_workbook(io.BytesIO(content))
    sheet = workbook["Productos"]
    headers = [str(cell.value) for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    assert headers[:5] == [
        "Nombre",
        "Cantidad inicial",
        "Estado de catálogo",
        "Precio de compra",
        "Precio de venta",
    ]
    assert headers[5] == "Aroma"
    sheet.append(["Vela lima", 4, "Activo", 1200, 3900, "Lima"])
    buffer = io.BytesIO()
    workbook.save(buffer)
    source = uploaded_asset(
        context,
        purpose=MediaAsset.Purpose.IMPORT_FILE,
        name="planilla-productos.xlsx",
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        content=buffer.getvalue(),
    )
    job = start_inventory_import(context=context, asset_id=source.public_id)
    analyse_import_job(job.public_id)
    mapping = mapping_from_import_headers(context, headers)
    preview_inventory_import(context=context, import_id=job.public_id, mapping=mapping)
    confirm_inventory_import(
        context=context,
        import_id=job.public_id,
        idempotency_key="confirm-template-1",
    )
    completed = process_import_job(job.public_id)
    assert completed.status == InventoryImport.Status.SUCCEEDED
    product = Product.objects.get(inventory=context.inventory, name="Vela lima")
    assert product.catalog_status == Product.CatalogStatus.ACTIVE
    assert product.purchase_price == 1200
    assert product.sale_price == 3900
    assert product.extra_attributes["aroma"] == "Lima"
    assert StockMovement.objects.filter(product=product, quantity=4).exists()


def test_export_is_idempotent_private_and_contains_filtered_inventory() -> None:
    fake_object_storage.clear()
    context = context_for("export-owner@example.com")
    create_product(context=context, name="Vela", initial_quantity=3)
    create_product(context=context, name="Taza", initial_quantity=2)

    first, replayed = start_inventory_export(
        context=context,
        file_format=InventoryExport.FileFormat.CSV,
        filters={"search": "Vela"},
        idempotency_key="export-vela-1",
    )
    second, replayed_second = start_inventory_export(
        context=context,
        file_format=InventoryExport.FileFormat.CSV,
        filters={"search": "Vela"},
        idempotency_key="export-vela-1",
    )
    assert replayed is False
    assert replayed_second is True
    assert first.pk == second.pk

    completed = process_export_job(first.public_id)
    assert completed.status == InventoryExport.Status.SUCCEEDED
    assert completed.row_count == 1
    assert completed.file_asset is not None
    content = fake_object_storage.read_bytes(key=completed.file_asset.object_key)
    assert b"Vela" in content
    assert b"Taza" not in content
    url = inventory_export_download_url(context=context, job=completed)
    assert url is not None
    assert f"/api/v1/media/{completed.file_asset.public_id}/content" in url
    assert "r2.invalid" not in url


def test_failed_export_retries_the_same_durable_job() -> None:
    context = context_for("export-retry@example.com")
    job = InventoryExport.objects.create(
        organisation=context.organisation,
        inventory=context.inventory,
        created_by=context.user,
        file_format=InventoryExport.FileFormat.CSV,
        status=InventoryExport.Status.FAILED,
        error_code="STORAGE_NOT_CONFIGURED",
    )

    retried = retry_inventory_export(context=context, export_id=job.public_id)

    assert retried.pk == job.pk
    assert retried.status == InventoryExport.Status.QUEUED
    assert retried.error_code == ""


def test_stock_alerts_are_deduplicated_resolved_and_queries_have_no_side_effects() -> None:
    context = context_for("alerts-owner@example.com")
    product = create_product(context=context, name="Vela", initial_quantity=5).product

    active = active_stock_alerts(context)
    assert len(active) == 1
    assert active[0].alert_type == InventoryAlert.AlertType.LOW_STOCK
    first_alert_id = active[0].id

    record_stock_movement(
        context=context,
        product_id=product.public_id,
        movement_type=StockMovement.MovementType.EXIT,
        quantity=1,
        idempotency_key="alert-exit-1",
    )
    active = active_stock_alerts(context)
    assert len(active) == 1
    assert active[0].id == first_alert_id
    assert active[0].available_quantity == 4

    record_stock_movement(
        context=context,
        product_id=product.public_id,
        movement_type=StockMovement.MovementType.EXIT,
        quantity=4,
        idempotency_key="alert-exit-2",
    )
    active = active_stock_alerts(context)
    assert len(active) == 1
    assert active[0].alert_type == InventoryAlert.AlertType.OUT_OF_STOCK
    assert InventoryAlert.objects.filter(
        product=product,
        alert_type=InventoryAlert.AlertType.LOW_STOCK,
        status=InventoryAlert.Status.RESOLVED,
    ).exists()

    count_before = InventoryAlert.objects.count()
    inventory_summary(context)
    active_stock_alerts(context)
    inventory_summary(context)
    assert InventoryAlert.objects.count() == count_before


def test_inventory_and_product_thresholds_recompute_alerts() -> None:
    context = context_for("threshold-owner@example.com")
    product = create_product(context=context, name="Taza", initial_quantity=6).product
    assert active_stock_alerts(context) == []

    update_inventory_low_stock_threshold(context=context, threshold=6)
    assert active_stock_alerts(context)[0].alert_type == InventoryAlert.AlertType.LOW_STOCK

    update_product_low_stock_threshold(
        context=context,
        product_id=product.public_id,
        threshold=0,
    )
    assert active_stock_alerts(context) == []

    update_product_low_stock_threshold(
        context=context,
        product_id=product.public_id,
        threshold=None,
    )
    assert active_stock_alerts(context)[0].threshold == 6


def test_foreign_jobs_are_not_visible() -> None:
    context = context_for("jobs-owner@example.com")
    foreign = context_for("jobs-foreign@example.com")
    source = uploaded_asset(
        context,
        purpose=MediaAsset.Purpose.IMPORT_FILE,
        name="products.csv",
        content_type="text/csv",
        content=b"Nombre\nVela\n",
    )
    job = start_inventory_import(context=context, asset_id=source.public_id)

    with pytest.raises(DomainError):
        preview_inventory_import(
            context=foreign,
            import_id=uuid.UUID(str(job.public_id)),
            mapping={"name": "Nombre"},
        )
