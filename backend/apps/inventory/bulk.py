"""Durable CSV/XLSX import and export jobs.

Every imported row has its own deterministic idempotency key and transaction.
One bad row therefore cannot roll back valid rows, and replaying a worker after
a crash never duplicates products or initial stock movements.
"""

from __future__ import annotations

import csv
import io
import logging
import uuid
from collections.abc import Mapping, Sequence
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from decimal import Decimal
from typing import Any
from zipfile import BadZipFile

from django.db import transaction
from django.utils import timezone
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.worksheet.datavalidation import DataValidation

from apps.audit.idempotency import execute_idempotent
from apps.audit.services import record_audit_event
from apps.media_assets.models import MediaAsset
from apps.media_assets.services import create_generated_asset, private_download_url
from apps.media_assets.storage import get_object_storage
from apps.organisations.selectors import TenantContext, resolve_tenant_context
from tenda.errors import DomainError, ResourceNotFound

from .custom_fields import validate_extra_attributes
from .models import InventoryExport, InventoryImport, Product
from .selectors import ProductFilters, active_custom_fields, filtered_products
from .services import MAX_MOVEMENT_QUANTITY, clean_reference_price, create_product

MAX_IMPORT_ROWS = 50_000
MAX_STORED_ERRORS = 1_000
PREVIEW_ROWS = 20
logger = logging.getLogger(__name__)

CORE_DESTINATIONS = frozenset(
    {
        "name",
        "initialQuantity",
        "catalogStatus",
        "purchasePrice",
        "salePrice",
    }
)

CATALOG_STATUS_ALIASES = {
    "active": Product.CatalogStatus.ACTIVE,
    "activo": Product.CatalogStatus.ACTIVE,
    "inactive": Product.CatalogStatus.INACTIVE,
    "inactivo": Product.CatalogStatus.INACTIVE,
}

IMPORT_TEMPLATE_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
IMPORT_TEMPLATE_FILENAME = "planilla-productos.xlsx"


@dataclass(frozen=True, slots=True)
class ImportColumn:
    destination: str
    header: str
    required: bool = False


IMPORT_CORE_COLUMNS: tuple[ImportColumn, ...] = (
    ImportColumn("name", "Nombre", True),
    ImportColumn("initialQuantity", "Cantidad inicial"),
    ImportColumn("catalogStatus", "Estado de catálogo"),
    ImportColumn("purchasePrice", "Precio de compra"),
    ImportColumn("salePrice", "Precio de venta"),
)


def import_template_columns(context: TenantContext) -> tuple[ImportColumn, ...]:
    """Stable Excel headers: core fields first, then each active custom column."""

    used = {column.header.casefold() for column in IMPORT_CORE_COLUMNS}
    columns = list(IMPORT_CORE_COLUMNS)
    for field in active_custom_fields(context):
        header = field.label.strip()
        if header.casefold() in used:
            header = f"{field.label.strip()} ({field.key})"
        used.add(header.casefold())
        columns.append(ImportColumn(f"extraAttributes.{field.key}", header))
    return tuple(columns)


def mapping_from_import_headers(
    context: TenantContext,
    headers: Sequence[str],
) -> dict[str, str]:
    folded: dict[str, str] = {}
    for header in headers:
        cleaned = str(header).strip()
        if cleaned:
            folded[cleaned.casefold()] = cleaned
    mapping: dict[str, str] = {}
    for column in import_template_columns(context):
        source = folded.get(column.header.casefold())
        if source:
            mapping[column.destination] = source
    return mapping


def build_import_template(context: TenantContext) -> tuple[str, str, bytes]:
    columns = import_template_columns(context)
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Productos"
    headers = [column.header for column in columns]
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    sheet.freeze_panes = "A2"
    last_column = get_column_letter(max(len(columns), 1))
    sheet.auto_filter.ref = f"A1:{last_column}1"
    for index, column in enumerate(columns, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = min(
            28,
            max(16, len(column.header) + 4),
        )

    status_index = next(
        (
            index
            for index, column in enumerate(columns, start=1)
            if column.destination == "catalogStatus"
        ),
        None,
    )
    if status_index is not None:
        letter = get_column_letter(status_index)
        validation = DataValidation(
            type="list",
            formula1='"Activo,Inactivo,active,inactive"',
            allow_blank=True,
        )
        validation.promptTitle = "Estado de catálogo"
        validation.prompt = "Usa Activo o Inactivo."
        sheet.add_data_validation(validation)
        validation.add(f"{letter}2:{letter}5000")

    instructions = workbook.create_sheet("Instrucciones")
    instructions["A1"] = "Cómo usar esta planilla"
    instructions["A1"].font = Font(bold=True)
    for index, line in enumerate(
        (
            "No cambies ni borres la fila de encabezados.",
            "Completa una fila por producto. El nombre es obligatorio y único en el inventario.",
            "Cantidad inicial es un entero de 0 o más. Si la dejas vacía, queda en 0.",
            "Estado de catálogo: Activo o Inactivo.",
            "Los precios van en pesos chilenos enteros, sin símbolo ni decimales.",
            "Las columnas extra son las que definiste en tu inventario.",
            "Cuando termines, súbela en Tenda web: Inventario → Importar.",
        ),
        start=3,
    ):
        instructions[f"A{index}"] = line
    instructions.column_dimensions["A"].width = 96

    buffer = io.BytesIO()
    workbook.save(buffer)
    return IMPORT_TEMPLATE_FILENAME, IMPORT_TEMPLATE_CONTENT_TYPE, buffer.getvalue()


@dataclass(frozen=True, slots=True)
class ImportFile:
    headers: tuple[str, ...]
    rows: tuple[dict[str, Any], ...]


def _json_cell(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return str(value)
    if isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def _normalise_headers(values: Sequence[Any]) -> tuple[str, ...]:
    headers = tuple(str(_json_cell(value)).strip() for value in values)
    if not headers or not any(headers):
        raise DomainError("INVALID_IMPORT_FILE", "El archivo no contiene encabezados.")
    if any(not header for header in headers):
        raise DomainError(
            "INVALID_IMPORT_FILE",
            "Todos los encabezados deben tener un nombre.",
        )
    folded = [header.casefold() for header in headers]
    if len(folded) != len(set(folded)):
        raise DomainError(
            "INVALID_IMPORT_FILE",
            "Los encabezados del archivo no pueden repetirse.",
        )
    return headers


def _read_csv(content: bytes) -> ImportFile:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise DomainError(
            "INVALID_IMPORT_FILE",
            "El CSV debe estar codificado en UTF-8.",
        ) from exc
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    reader = csv.reader(io.StringIO(text), dialect)
    try:
        headers = _normalise_headers(next(reader))
    except StopIteration as exc:
        raise DomainError("INVALID_IMPORT_FILE", "El archivo está vacío.") from exc
    rows: list[dict[str, Any]] = []
    for values in reader:
        if not any(str(value).strip() for value in values):
            continue
        if len(rows) >= MAX_IMPORT_ROWS:
            raise DomainError(
                "IMPORT_ROW_LIMIT",
                f"El archivo admite hasta {MAX_IMPORT_ROWS} filas.",
            )
        padded = [*values, *([""] * max(0, len(headers) - len(values)))]
        rows.append(
            {
                header: _json_cell(padded[index]) if index < len(padded) else ""
                for index, header in enumerate(headers)
            }
        )
    return ImportFile(headers=headers, rows=tuple(rows))


def _read_xlsx(content: bytes) -> ImportFile:
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        sheet = workbook.active
        if sheet is None:
            raise ValueError("workbook has no active sheet")
        iterator = sheet.iter_rows(values_only=True)
        headers = _normalise_headers(next(iterator))
    except (OSError, ValueError, StopIteration, BadZipFile, InvalidFileException) as exc:
        raise DomainError(
            "INVALID_IMPORT_FILE",
            "No pudimos leer el archivo XLSX.",
        ) from exc
    rows: list[dict[str, Any]] = []
    for values in iterator:
        if not any(value not in (None, "") for value in values):
            continue
        if len(rows) >= MAX_IMPORT_ROWS:
            raise DomainError(
                "IMPORT_ROW_LIMIT",
                f"El archivo admite hasta {MAX_IMPORT_ROWS} filas.",
            )
        rows.append(
            {
                header: _json_cell(values[index]) if index < len(values) else ""
                for index, header in enumerate(headers)
            }
        )
    workbook.close()
    return ImportFile(headers=headers, rows=tuple(rows))


def read_import_file(asset: MediaAsset) -> ImportFile:
    content = get_object_storage().read_bytes(key=asset.object_key)
    if not content:
        raise DomainError("INVALID_IMPORT_FILE", "El archivo está vacío.")
    if asset.content_type == "text/csv":
        return _read_csv(content)
    if asset.content_type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet":
        return _read_xlsx(content)
    raise DomainError("INVALID_IMPORT_FILE", "El formato del archivo no está permitido.")


def _context_for_import(job: InventoryImport) -> TenantContext:
    return resolve_tenant_context(
        job.created_by,
        organisation_id=job.organisation.public_id,
        inventory_id=job.inventory.public_id,
    )


def _context_for_export(job: InventoryExport) -> TenantContext:
    return resolve_tenant_context(
        job.created_by,
        organisation_id=job.organisation.public_id,
        inventory_id=job.inventory.public_id,
    )


def validate_import_mapping(
    context: TenantContext,
    *,
    headers: Sequence[str],
    mapping: Mapping[str, Any],
) -> dict[str, str]:
    dynamic = {f"extraAttributes.{field.key}" for field in active_custom_fields(context)}
    allowed = CORE_DESTINATIONS | dynamic
    header_set = set(headers)
    clean: dict[str, str] = {}
    field_errors: dict[str, list[str]] = {}
    for raw_destination, raw_source in mapping.items():
        destination = str(raw_destination).strip()
        source = str(raw_source).strip()
        if destination not in allowed:
            field_errors[f"mapping.{destination}"] = [
                "Ese campo no existe en el esquema del inventario."
            ]
            continue
        if source not in header_set:
            field_errors[f"mapping.{destination}"] = [
                "La columna de origen no existe en el archivo."
            ]
            continue
        clean[destination] = source
    if "name" not in clean:
        field_errors["mapping.name"] = ["Asocia una columna con el nombre del producto."]
    sources = list(clean.values())
    if len(sources) != len(set(sources)):
        field_errors["mapping"] = ["Cada columna del archivo puede usarse una sola vez."]
    if field_errors:
        raise DomainError(
            "INVALID_IMPORT_MAPPING",
            "Revisa el mapeo de columnas.",
            field_errors=field_errors,
        )
    return clean


def _row_value(row: Mapping[str, Any], mapping: Mapping[str, str], key: str) -> Any:
    source = mapping.get(key)
    return row.get(source, "") if source else ""


def normalise_import_row(
    context: TenantContext,
    *,
    row: Mapping[str, Any],
    mapping: Mapping[str, str],
) -> dict[str, Any]:
    name = " ".join(str(_row_value(row, mapping, "name")).split())
    if not name:
        raise DomainError(
            "VALIDATION_ERROR",
            "Revisa los datos ingresados.",
            field_errors={"name": ["Ingresa el nombre del producto."]},
        )
    if len(name) > 160:
        raise DomainError(
            "VALIDATION_ERROR",
            "Revisa los datos ingresados.",
            field_errors={"name": ["Usa un nombre de hasta 160 caracteres."]},
        )
    raw_status = str(_row_value(row, mapping, "catalogStatus")).strip()
    if not raw_status:
        catalog_status = Product.CatalogStatus.ACTIVE
    else:
        catalog_status = CATALOG_STATUS_ALIASES.get(raw_status.casefold())
        if catalog_status is None:
            raise DomainError(
                "VALIDATION_ERROR",
                "Revisa los datos ingresados.",
                field_errors={"catalogStatus": ["Usa Activo o Inactivo."]},
            )
    initial_raw = _row_value(row, mapping, "initialQuantity")
    if initial_raw in (None, ""):
        initial_quantity = 0
    else:
        try:
            initial_quantity = int(str(initial_raw).strip())
        except ValueError as exc:
            raise DomainError(
                "VALIDATION_ERROR",
                "Revisa los datos ingresados.",
                field_errors={"initialQuantity": ["Ingresa una cantidad entera de 0 o más."]},
            ) from exc
        if initial_quantity < 0 or initial_quantity > MAX_MOVEMENT_QUANTITY:
            raise DomainError(
                "VALIDATION_ERROR",
                "Revisa los datos ingresados.",
                field_errors={"initialQuantity": ["Ingresa una cantidad entera de 0 o más."]},
            )
    attributes = {
        destination.removeprefix("extraAttributes."): _row_value(row, mapping, destination)
        for destination in mapping
        if destination.startswith("extraAttributes.")
    }
    return {
        "name": name,
        "catalog_status": catalog_status,
        "purchase_price": clean_reference_price(
            _row_value(row, mapping, "purchasePrice"),
            field="purchasePrice",
        ),
        "sale_price": clean_reference_price(
            _row_value(row, mapping, "salePrice"),
            field="salePrice",
        ),
        "initial_quantity": initial_quantity,
        "extra_attributes": validate_extra_attributes(
            active_custom_fields(context),
            attributes,
        ),
    }


def _error_payload(row_number: int, error: DomainError) -> dict[str, Any]:
    return {
        "row": row_number,
        "code": error.code,
        "message": error.message,
        "fieldErrors": error.field_errors,
    }


@transaction.atomic
def start_inventory_import(
    *,
    context: TenantContext,
    asset_id: uuid.UUID,
    correlation_id: str = "",
) -> InventoryImport:
    asset = (
        MediaAsset.objects.select_for_update()
        .filter(
            public_id=asset_id,
            organisation=context.organisation,
            purpose=MediaAsset.Purpose.IMPORT_FILE,
            status=MediaAsset.Status.READY,
        )
        .first()
    )
    if asset is None:
        raise ResourceNotFound()
    existing = InventoryImport.objects.filter(
        inventory=context.inventory,
        source_asset=asset,
    ).first()
    if existing is not None:
        return existing
    job = InventoryImport.objects.create(
        organisation=context.organisation,
        inventory=context.inventory,
        source_asset=asset,
        created_by=context.user,
    )
    record_audit_event(
        action="inventory.import_started",
        organisation=context.organisation,
        actor=context.user,
        object_type="inventory.import",
        object_public_id=str(job.public_id),
        correlation_id=correlation_id,
        metadata={"fileType": asset.content_type, "size": asset.actual_size},
    )

    def enqueue() -> None:
        from .tasks import analyse_inventory_import

        analyse_inventory_import.delay(str(job.public_id))

    transaction.on_commit(enqueue)
    return job


def analyse_import_job(public_id: uuid.UUID) -> InventoryImport:
    job = (
        InventoryImport.objects.select_related(
            "source_asset",
            "organisation",
            "inventory",
            "created_by",
        )
        .filter(public_id=public_id)
        .first()
    )
    if job is None:
        raise ResourceNotFound()
    if job.status != InventoryImport.Status.ANALYSING:
        return job
    try:
        parsed = read_import_file(job.source_asset)
        job.headers = list(parsed.headers)
        job.preview_rows = list(parsed.rows[:PREVIEW_ROWS])
        job.total_rows = len(parsed.rows)
        job.progress = 100
        job.status = InventoryImport.Status.AWAITING_MAPPING
        job.error_code = ""
        job.save(
            update_fields=(
                "headers",
                "preview_rows",
                "total_rows",
                "progress",
                "status",
                "error_code",
                "updated_at",
            )
        )
    except DomainError as exc:
        job.status = InventoryImport.Status.FAILED
        job.error_code = exc.code
        job.completed_at = timezone.now()
        job.save(update_fields=("status", "error_code", "completed_at", "updated_at"))
    except Exception:
        logger.exception(
            "Unexpected inventory import analysis failure",
            extra={"job": str(public_id)},
        )
        job.status = InventoryImport.Status.FAILED
        job.error_code = "UNEXPECTED_ERROR"
        job.completed_at = timezone.now()
        job.save(update_fields=("status", "error_code", "completed_at", "updated_at"))
    return job


@transaction.atomic
def preview_inventory_import(
    *,
    context: TenantContext,
    import_id: uuid.UUID,
    mapping: Mapping[str, Any],
) -> InventoryImport:
    job = (
        InventoryImport.objects.select_for_update()
        .select_related("source_asset")
        .filter(
            public_id=import_id,
            inventory=context.inventory,
            organisation=context.organisation,
        )
        .first()
    )
    if job is None:
        raise ResourceNotFound()
    if job.status != InventoryImport.Status.AWAITING_MAPPING:
        raise DomainError(
            "IMPORT_NOT_READY",
            "La importación todavía no está lista para mapear.",
            status=409,
            retryable=job.status == InventoryImport.Status.ANALYSING,
        )
    clean_mapping = validate_import_mapping(
        context,
        headers=job.headers,
        mapping=mapping,
    )
    parsed = read_import_file(job.source_asset)
    errors: list[dict[str, Any]] = []
    for index, row in enumerate(parsed.rows[:PREVIEW_ROWS], start=2):
        try:
            normalise_import_row(context, row=row, mapping=clean_mapping)
        except DomainError as exc:
            errors.append(_error_payload(index, exc))
    job.mapping = clean_mapping
    job.row_errors = errors
    job.preview_rows = list(parsed.rows[:PREVIEW_ROWS])
    job.save(update_fields=("mapping", "row_errors", "preview_rows", "updated_at"))
    return job


def confirm_inventory_import(
    *,
    context: TenantContext,
    import_id: uuid.UUID,
    idempotency_key: str,
    correlation_id: str = "",
) -> tuple[InventoryImport, bool]:
    job = InventoryImport.objects.filter(
        public_id=import_id,
        inventory=context.inventory,
        organisation=context.organisation,
    ).first()
    if job is None:
        raise ResourceNotFound()

    def command() -> tuple[dict[str, Any], int]:
        locked = InventoryImport.objects.select_for_update().get(pk=job.pk)
        if locked.status != InventoryImport.Status.AWAITING_MAPPING or not locked.mapping:
            raise DomainError(
                "IMPORT_NOT_READY",
                "Revisa el mapeo antes de confirmar la importación.",
                status=409,
            )
        locked.status = InventoryImport.Status.QUEUED
        locked.progress = 0
        locked.processed_rows = 0
        locked.created_count = 0
        locked.error_count = 0
        locked.error_code = ""
        locked.started_at = None
        locked.completed_at = None
        locked.save(
            update_fields=(
                "status",
                "progress",
                "processed_rows",
                "created_count",
                "error_count",
                "error_code",
                "started_at",
                "completed_at",
                "updated_at",
            )
        )

        def enqueue() -> None:
            from .tasks import process_inventory_import

            process_inventory_import.delay(str(locked.public_id))

        transaction.on_commit(enqueue)
        return {"importId": str(locked.public_id)}, 202

    outcome = execute_idempotent(
        context=context,
        scope="inventory.confirm_import",
        key=idempotency_key,
        request_payload={"importId": str(import_id), "mapping": job.mapping},
        command=command,
    )
    refreshed = InventoryImport.objects.get(public_id=uuid.UUID(str(outcome.payload["importId"])))
    record_audit_event(
        action="inventory.import_confirmed",
        organisation=context.organisation,
        actor=context.user,
        object_type="inventory.import",
        object_public_id=str(refreshed.public_id),
        correlation_id=correlation_id,
        metadata={"replayed": outcome.replayed, "rows": refreshed.total_rows},
    )
    return refreshed, outcome.replayed


def _import_report(errors: Sequence[Mapping[str, Any]], total: int, created: int) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(("fila", "estado", "codigo", "detalle"))
    failed_rows = {int(error["row"]): error for error in errors}
    for row_number in range(2, total + 2):
        error = failed_rows.get(row_number)
        if error:
            writer.writerow(
                (
                    row_number,
                    "error",
                    error.get("code", ""),
                    error.get("message", ""),
                )
            )
        else:
            writer.writerow((row_number, "creado", "", ""))
    writer.writerow(("", "resumen", "", f"{created} creados; {len(errors)} con error"))
    return buffer.getvalue().encode("utf-8-sig")


def process_import_job(public_id: uuid.UUID) -> InventoryImport:
    job = (
        InventoryImport.objects.select_related(
            "source_asset",
            "organisation",
            "inventory",
            "created_by",
        )
        .filter(public_id=public_id)
        .first()
    )
    if job is None:
        raise ResourceNotFound()
    if job.status in {
        InventoryImport.Status.SUCCEEDED,
        InventoryImport.Status.COMPLETED_WITH_ERRORS,
    }:
        return job
    if job.status not in {
        InventoryImport.Status.QUEUED,
        InventoryImport.Status.PROCESSING,
        InventoryImport.Status.FAILED,
    }:
        return job
    context = _context_for_import(job)
    job.status = InventoryImport.Status.PROCESSING
    job.started_at = job.started_at or timezone.now()
    job.completed_at = None
    job.error_code = ""
    job.save(update_fields=("status", "started_at", "completed_at", "error_code", "updated_at"))
    try:
        parsed = read_import_file(job.source_asset)
        mapping = validate_import_mapping(
            context,
            headers=parsed.headers,
            mapping=job.mapping,
        )
        errors: list[dict[str, Any]] = []
        created = 0
        total = len(parsed.rows)
        for offset, row in enumerate(parsed.rows, start=1):
            row_number = offset + 1
            try:
                values = normalise_import_row(context, row=row, mapping=mapping)
                create_product(
                    context=context,
                    **values,
                    reason="Importación de inventario",
                    idempotency_key=f"import:{job.public_id}:row:{row_number}",
                    correlation_id=f"import:{job.public_id}",
                )
                created += 1
            except DomainError as exc:
                errors.append(_error_payload(row_number, exc))
            if offset % 25 == 0 or offset == total:
                job.processed_rows = offset
                job.created_count = created
                job.error_count = len(errors)
                job.progress = int(offset * 100 / max(total, 1))
                job.save(
                    update_fields=(
                        "processed_rows",
                        "created_count",
                        "error_count",
                        "progress",
                        "updated_at",
                    )
                )
        report = create_generated_asset(
            context=context,
            purpose=MediaAsset.Purpose.IMPORT_REPORT,
            original_name=f"reporte-importacion-{job.public_id}.csv",
            content_type="text/csv",
            content=_import_report(errors, total, created),
        )
        job.report_asset = report
        job.row_errors = errors[:MAX_STORED_ERRORS]
        job.total_rows = total
        job.processed_rows = total
        job.created_count = created
        job.error_count = len(errors)
        job.progress = 100
        job.status = (
            InventoryImport.Status.COMPLETED_WITH_ERRORS
            if errors
            else InventoryImport.Status.SUCCEEDED
        )
        job.completed_at = timezone.now()
        job.save(
            update_fields=(
                "report_asset",
                "row_errors",
                "total_rows",
                "processed_rows",
                "created_count",
                "error_count",
                "progress",
                "status",
                "completed_at",
                "updated_at",
            )
        )
        record_audit_event(
            action="inventory.import_completed",
            organisation=context.organisation,
            actor=context.user,
            object_type="inventory.import",
            object_public_id=str(job.public_id),
            correlation_id=f"import:{job.public_id}",
            metadata={"created": created, "errors": len(errors), "total": total},
        )
    except DomainError as exc:
        job.status = InventoryImport.Status.FAILED
        job.error_code = exc.code
        job.completed_at = timezone.now()
        job.save(update_fields=("status", "error_code", "completed_at", "updated_at"))
    except Exception:
        logger.exception("Unexpected inventory import failure", extra={"job": str(public_id)})
        job.status = InventoryImport.Status.FAILED
        job.error_code = "UNEXPECTED_ERROR"
        job.completed_at = timezone.now()
        job.save(update_fields=("status", "error_code", "completed_at", "updated_at"))
    return job


@transaction.atomic
def retry_inventory_import(
    *,
    context: TenantContext,
    import_id: uuid.UUID,
) -> InventoryImport:
    job = (
        InventoryImport.objects.select_for_update()
        .filter(
            public_id=import_id,
            inventory=context.inventory,
            organisation=context.organisation,
        )
        .first()
    )
    if job is None:
        raise ResourceNotFound()
    if job.status not in {
        InventoryImport.Status.FAILED,
        InventoryImport.Status.COMPLETED_WITH_ERRORS,
    }:
        raise DomainError(
            "IMPORT_NOT_RETRYABLE",
            "Esta importación no necesita reintento.",
            status=409,
        )
    job.status = InventoryImport.Status.QUEUED
    job.progress = 0
    job.error_code = ""
    job.completed_at = None
    job.save(update_fields=("status", "progress", "error_code", "completed_at", "updated_at"))

    def enqueue() -> None:
        from .tasks import process_inventory_import

        process_inventory_import.delay(str(job.public_id))

    transaction.on_commit(enqueue)
    return job


def inventory_import_for_context(
    context: TenantContext,
    public_id: uuid.UUID,
) -> InventoryImport:
    job = (
        InventoryImport.objects.filter(
            public_id=public_id,
            inventory=context.inventory,
            organisation=context.organisation,
        )
        .select_related("source_asset", "report_asset")
        .first()
    )
    if job is None:
        raise ResourceNotFound()
    return job


def inventory_imports_for_context(context: TenantContext) -> list[InventoryImport]:
    return list(
        InventoryImport.objects.filter(
            inventory=context.inventory,
            organisation=context.organisation,
        )
        .select_related("source_asset", "report_asset")
        .order_by("-created_at", "-id")[:50]
    )


def _filters_from_payload(payload: Mapping[str, Any]) -> ProductFilters:
    raw_attributes = payload.get("attributes")
    attributes: list[tuple[str, str]] = []
    if isinstance(raw_attributes, Sequence) and not isinstance(raw_attributes, str):
        for item in raw_attributes:
            if isinstance(item, Mapping):
                attributes.append((str(item.get("key", "")), str(item.get("value", ""))))
    return ProductFilters(
        search=str(payload.get("search") or ""),
        catalog_statuses=tuple(str(item) for item in (payload.get("catalogStatuses") or ())),
        stock_states=tuple(str(item) for item in (payload.get("stockStates") or ())),
        attributes=tuple(attributes),
        include_archived=bool(payload.get("includeArchived")),
    )


def start_inventory_export(
    *,
    context: TenantContext,
    file_format: str,
    filters: Mapping[str, Any] | None,
    idempotency_key: str,
    correlation_id: str = "",
) -> tuple[InventoryExport, bool]:
    if file_format not in InventoryExport.FileFormat.values:
        raise DomainError(
            "VALIDATION_ERROR",
            "Revisa los datos ingresados.",
            field_errors={"format": ["Selecciona CSV o XLSX."]},
        )
    clean_filters = dict(filters or {})
    # Execute the selector once before scheduling so invalid dynamic filters fail
    # in the request rather than several seconds later in a worker.
    filtered_products(context, _filters_from_payload(clean_filters)).exists()

    def command() -> tuple[dict[str, Any], int]:
        job = InventoryExport.objects.create(
            organisation=context.organisation,
            inventory=context.inventory,
            created_by=context.user,
            file_format=file_format,
            filters=clean_filters,
        )

        def enqueue() -> None:
            from .tasks import process_inventory_export

            process_inventory_export.delay(str(job.public_id))

        transaction.on_commit(enqueue)
        return {"exportId": str(job.public_id)}, 202

    outcome = execute_idempotent(
        context=context,
        scope="inventory.start_export",
        key=idempotency_key,
        request_payload={"format": file_format, "filters": clean_filters},
        command=command,
    )
    job = InventoryExport.objects.get(public_id=uuid.UUID(str(outcome.payload["exportId"])))
    record_audit_event(
        action="inventory.export_started",
        organisation=context.organisation,
        actor=context.user,
        object_type="inventory.export",
        object_public_id=str(job.public_id),
        correlation_id=correlation_id,
        metadata={"format": file_format, "replayed": outcome.replayed},
    )
    return job, outcome.replayed


@transaction.atomic
def retry_inventory_export(
    *,
    context: TenantContext,
    export_id: uuid.UUID,
) -> InventoryExport:
    job = (
        InventoryExport.objects.select_for_update()
        .filter(
            public_id=export_id,
            inventory=context.inventory,
            organisation=context.organisation,
        )
        .first()
    )
    if job is None:
        raise ResourceNotFound()
    if job.status != InventoryExport.Status.FAILED:
        raise DomainError(
            "EXPORT_NOT_RETRYABLE",
            "Esta exportación no necesita reintento.",
            status=409,
        )
    job.status = InventoryExport.Status.QUEUED
    job.progress = 0
    job.error_code = ""
    job.completed_at = None
    job.save(update_fields=("status", "progress", "error_code", "completed_at", "updated_at"))

    def enqueue() -> None:
        from .tasks import process_inventory_export

        process_inventory_export.delay(str(job.public_id))

    transaction.on_commit(enqueue)
    return job


def _export_rows(context: TenantContext, job: InventoryExport) -> tuple[list[str], list[list[Any]]]:
    fields = active_custom_fields(context)
    headers = [
        "Nombre",
        "Estado",
        "Stock total",
        "Reservado",
        "Disponible",
        "Precio compra",
        "Precio venta",
        *[field.label for field in fields],
    ]
    rows: list[list[Any]] = []
    products = filtered_products(context, _filters_from_payload(job.filters)).order_by("name", "id")
    for product in products:
        rows.append(
            [
                product.name,
                product.catalog_status,
                int(getattr(product, "on_hand", 0)),
                int(getattr(product, "reserved", 0)),
                int(getattr(product, "available", 0)),
                str(product.purchase_price) if product.purchase_price is not None else "",
                str(product.sale_price) if product.sale_price is not None else "",
                *[_json_cell((product.extra_attributes or {}).get(field.key)) for field in fields],
            ]
        )
    return headers, rows


def _csv_export(headers: Sequence[str], rows: Sequence[Sequence[Any]]) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(headers)
    writer.writerows(rows)
    return buffer.getvalue().encode("utf-8-sig")


def _xlsx_export(headers: Sequence[str], rows: Sequence[Sequence[Any]]) -> bytes:
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet(title="Inventario")
    sheet.append(list(headers))
    for row in rows:
        sheet.append(list(row))
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def process_export_job(public_id: uuid.UUID) -> InventoryExport:
    job = (
        InventoryExport.objects.select_related(
            "organisation",
            "inventory",
            "created_by",
            "file_asset",
        )
        .filter(public_id=public_id)
        .first()
    )
    if job is None:
        raise ResourceNotFound()
    if job.status == InventoryExport.Status.SUCCEEDED:
        return job
    context = _context_for_export(job)
    job.status = InventoryExport.Status.PROCESSING
    job.started_at = timezone.now()
    job.progress = 10
    job.error_code = ""
    job.save(update_fields=("status", "started_at", "progress", "error_code", "updated_at"))
    try:
        headers, rows = _export_rows(context, job)
        if job.file_format == InventoryExport.FileFormat.CSV:
            content = _csv_export(headers, rows)
            content_type = "text/csv"
            extension = "csv"
        else:
            content = _xlsx_export(headers, rows)
            content_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            extension = "xlsx"
        asset = create_generated_asset(
            context=context,
            purpose=MediaAsset.Purpose.INVENTORY_EXPORT,
            original_name=f"inventario-{job.created_at.date()}.{extension}",
            content_type=content_type,
            content=content,
        )
        job.file_asset = asset
        job.row_count = len(rows)
        job.progress = 100
        job.status = InventoryExport.Status.SUCCEEDED
        job.expires_at = timezone.now() + timedelta(hours=24)
        job.completed_at = timezone.now()
        job.save(
            update_fields=(
                "file_asset",
                "row_count",
                "progress",
                "status",
                "expires_at",
                "completed_at",
                "updated_at",
            )
        )
    except DomainError as exc:
        job.status = InventoryExport.Status.FAILED
        job.error_code = exc.code
        job.completed_at = timezone.now()
        job.save(update_fields=("status", "error_code", "completed_at", "updated_at"))
    except Exception:
        logger.exception("Unexpected inventory export failure", extra={"job": str(public_id)})
        job.status = InventoryExport.Status.FAILED
        job.error_code = "UNEXPECTED_ERROR"
        job.completed_at = timezone.now()
        job.save(update_fields=("status", "error_code", "completed_at", "updated_at"))
    return job


def inventory_export_for_context(
    context: TenantContext,
    public_id: uuid.UUID,
) -> InventoryExport:
    job = (
        InventoryExport.objects.filter(
            public_id=public_id,
            inventory=context.inventory,
            organisation=context.organisation,
        )
        .select_related("file_asset")
        .first()
    )
    if job is None:
        raise ResourceNotFound()
    return job


def inventory_exports_for_context(context: TenantContext) -> list[InventoryExport]:
    return list(
        InventoryExport.objects.filter(
            inventory=context.inventory,
            organisation=context.organisation,
        )
        .select_related("file_asset")
        .order_by("-created_at", "-id")[:50]
    )


def inventory_export_download_url(
    *,
    context: TenantContext,
    job: InventoryExport,
) -> str | None:
    if (
        job.status != InventoryExport.Status.SUCCEEDED
        or job.file_asset is None
        or job.expires_at is None
        or job.expires_at <= timezone.now()
    ):
        return None
    return private_download_url(context=context, public_id=job.file_asset.public_id)


def inventory_import_report_url(
    *,
    context: TenantContext,
    job: InventoryImport,
) -> str | None:
    if job.report_asset is None:
        return None
    return private_download_url(context=context, public_id=job.report_asset.public_id)
